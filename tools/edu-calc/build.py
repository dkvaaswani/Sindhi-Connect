"""Build the Excel version of the Children's Future Education Fund calculator (method v2).

    py tools/edu-calc/build.py

Reads   data/education/education-costs.json                       (the master dataset)
Writes  frontend/downloads/Children_Education_Fund_Calculator.xlsx (formula-driven workbook, no macros)

The workbook follows the same rules as frontend/edu-calc/calc-engine.js (see
docs/education-calculator/METHODOLOGY.md):
  - published fee records are kept in the workbook and averaged with live formulas, so a changed
    fee increase re-averages them exactly as the website does;
  - planning estimates, pre-stage fees, living costs, visas and travel are resolved here with the
    same rules and stored as tables;
  - each child's yearly costs, the amount needed now, the required yearly saving and the yearly fund
    roll-forward are all live formulas.

Needs Python 3 with openpyxl. Then run tests/edu-calc/excel-parity.ps1 (needs Excel) to recalculate
the workbook, store its values, and check it against the web engine.
"""
import json
import math
import os
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
DATA = os.path.join(ROOT, "data", "education", "education-costs.json")
OUT_XLSX = os.path.join(ROOT, "frontend", "downloads", "Children_Education_Fund_Calculator.xlsx")

KIDS = 4
ROWS = 60           # plan years per child (age 0 + entry age up to 45 + course length)
STUDY = 10          # study years per course (longest course is 8 years)
CHART_ROWS = 30

# ---- styles (Sindhi Connect "Midnight & Gold") ----
INK, GOLD, LINE = "0E1016", "E9A825", "DDD8CC"
F = "Arial"
font = lambda **k: Font(name=F, size=k.pop("size", 10), **k)
TITLE = font(size=16, bold=True, color=INK)
H2 = font(size=12, bold=True, color=INK)
HDR = font(bold=True, color="FFFFFF")
HDR_FILL = PatternFill("solid", fgColor=INK)
INPUT_FILL = PatternFill("solid", fgColor="FFF4D6")
CALC_FILL = PatternFill("solid", fgColor="F2F2F2")
GOLD_FILL = PatternFill("solid", fgColor=GOLD)
MUTED = font(color="565B66", italic=True, size=9)
BOLD = font(bold=True)
BIG = font(size=14, bold=True, color=INK)
thin = Side(style="thin", color=LINE)
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")
NUM = '#,##0;-#,##0;"-"'
PCT = '0.0%'

# Cost lines, in the engine's order: key, label, increase type, timing, scholarship applies
ITEMS = [
    ("tuition", "Tuition", "tuition", "annual", True),
    ("accommodation", "Accommodation", "living", "annual", True),
    ("food", "Food", "living", "annual", True),
    ("transport", "Local transport", "living", "annual", True),
    ("healthInsurance", "Health insurance", "living", "annual", True),
    ("books", "Books and supplies", "living", "annual", True),
    ("otherFees", "Other university and course fees", "tuition", "annual", True),
    ("admission", "Admission and one-time fees", "tuition", "once", True),
    ("visaApplication", "Visa and application", "living", "once", False),
    ("travelRelocation", "Travel and relocation", "living", "once", False),
]
PRE_ITEMS = [("preTuition", "Pre-stage tuition (e.g. pre-medical degree)"),
             ("preOtherFees", "Pre-stage other fees")]
LIVING_KEYS = ["accommodation", "food", "transport", "healthInsurance", "books"]
GROUPS = [("Tuition", ["tuition"]), ("Accommodation", ["accommodation"]), ("Food", ["food"]),
          ("Local transport", ["transport"]), ("Health insurance", ["healthInsurance"]),
          ("Books and supplies", ["books"]), ("Other university and course fees", ["otherFees", "admission"]),
          ("Visa, application, travel and relocation", ["visaApplication", "travelRelocation"])]
STATUS_LABEL = ('IF({s}="verified","Published figure",IF({s}="estimated","Estimated — please review",'
                'IF({s}="notNeeded","Not needed",IF({s}="override","Your figure",IF({s}="missing","No figure yet — enter your own","")))))')


def q(sheet):
    return "'" + sheet + "'"


def put(ws, ref, value, f=None, fill=None, fmt=None, wrap=False, lock=None, border=False, align=None):
    c = ws[ref]
    c.value = value
    c.font = f or font()
    if fill:
        c.fill = fill
    if fmt:
        c.number_format = fmt
    if wrap:
        c.alignment = WRAP
    if align:
        c.alignment = align
    if lock is not None:
        c.protection = Protection(locked=lock)
    if border:
        c.border = BOX
    return c


def header_row(ws, row, col, labels, widths=None):
    for i, lab in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=lab)
        c.font, c.fill, c.alignment, c.border = HDR, HDR_FILL, Alignment(wrap_text=True, vertical="center"), BOX
        if widths:
            ws.column_dimensions[L(col + i)].width = widths[i]


def isnum(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool) and math.isfinite(v)


def num(v):
    return v if isnum(v) else 0


# =====================================================================
# Cost rules (same as calc-engine.js)
# =====================================================================
class Rules:
    def __init__(self, d):
        self.d = d
        self.plan = d["planStartYear"]
        self.rates = d["exchangeRates"]["rates"]

    def country(self, name):
        return next((c for c in self.d["countries"] if c["name"] == name), None)

    def qual(self, name):
        return next((x for x in self.d["qualifications"] if x["name"] == name), None)

    def convert(self, amount, frm, to):
        if frm == to or not amount:
            return amount
        return amount / self.rates[frm] * self.rates[to]

    def record_fee_status(self, r):
        return r.get("feeStatus") or self.d.get("feeStatusByCategory", {}).get(r.get("studentCategory")) or "international"

    def duration(self, country, qualification):
        d = next((x for x in self.d["courseDurations"] if x["country"] == country and x["qualification"] == qualification), None)
        if d:
            return {"years": d.get("years"), "available": isnum(d.get("years")) and d["years"] > 0,
                    "preStage": d.get("preStage"), "note": d.get("note", "")}
        recs = [r for r in self.d["records"] if r["country"] == country and r["qualification"] == qualification and isnum(r.get("durationYears"))]
        if recs:
            return {"years": sum(r["durationYears"] for r in recs) / len(recs), "available": True, "preStage": None, "note": ""}
        return {"years": None, "available": False, "preStage": None, "note": "No typical duration recorded."}

    def tuition_inflation(self, country, qualification, fee_status):
        q_ = self.qual(qualification)
        group = q_["group"] if q_ else None
        lst = [t for t in self.d.get("tuitionInflation", []) if t["country"] == country and
               (not t.get("feeStatus") or not fee_status or t.get("feeStatus") == fee_status)]
        best = lambda arr: sorted(arr, key=lambda t: 0 if t.get("feeStatus") else 1)[0] if arr else None
        hit = (best([t for t in lst if t.get("qualification") == qualification]) or
               best([t for t in lst if group and t.get("group") == group]) or
               best([t for t in lst if not t.get("group") and not t.get("qualification")]))
        if hit:
            return hit["rate"]
        info = self.country(country)
        return info["tuitionInflation"] if info else 0.05

    def used_records(self, country, qualification, fee_status):
        recs = [r for r in self.d["records"] if r["country"] == country and r["qualification"] == qualification and
                r.get("includeInAverage") and isnum(r.get("tuition")) and self.record_fee_status(r) == fee_status]
        if not recs:
            return None, []
        counts = {}
        for r in recs:
            counts[r["currency"]] = counts.get(r["currency"], 0) + 1
        ccy = sorted(counts.keys(), key=lambda c: -counts[c])[0]   # stable, like the engine
        return ccy, [r for r in recs if r["currency"] == ccy]

    def verified(self, country, qualification, fee_status, rate):
        ccy, same = self.used_records(country, qualification, fee_status)
        if not same:
            return None
        grow = lambda r: (1 + rate) ** (self.plan - (r.get("feeYearStart") or self.plan))
        mean = lambda vals: sum(vals) / len(vals)
        return {"status": "verified", "currency": ccy, "baseYear": self.plan, "count": len(same),
                "tuition": mean([r["tuition"] * grow(r) for r in same]),
                "otherFees": mean([num(r.get("otherMandatoryAnnual")) * grow(r) for r in same]),
                "admission": mean([(num(r.get("oneTimeAdmission")) + num(r.get("otherOneTime"))) * grow(r) for r in same]),
                "note": "From one university's published fees." if len(same) == 1 else
                        "Average of %d universities' published fees." % len(same)}

    def derive(self, country, fee_status, rule):
        def ver(qn, fs):
            return self.verified(country, qn, fs, self.tuition_inflation(country, qn, fs))
        if rule.get("average") == "nonMedicalDegrees":
            info = self.country(country)
            ccy = info["currency"] if info else "USD"
            lst = [ver(x["name"], fee_status) for x in self.d["qualifications"]
                   if x["kind"] == "degree" and x["group"] != "medical" and x["name"] != "Other Qualification"]
            lst = [b for b in lst if b]
            if not lst:
                return None
            avg = lambda k: sum(self.convert(b[k], b["currency"], ccy) for b in lst) / len(lst)
            return {"currency": ccy, "baseYear": self.plan, "tuition": avg("tuition"), "otherFees": avg("otherFees"), "admission": avg("admission")}
        src = ver(rule["qualification"], rule.get("feeStatus") or fee_status)
        if not src:
            return None
        m = rule["multiplier"] if isnum(rule.get("multiplier")) else 1
        return {"currency": src["currency"], "baseYear": src["baseYear"], "tuition": src["tuition"] * m,
                "otherFees": src["otherFees"] * m, "admission": src["admission"] * m}

    def estimate(self, country, qualification, fee_status):
        """The engine's fee benchmark when no published record applies (estimate or missing)."""
        est = next((e for e in self.d.get("estimates", []) if e["country"] == country and e["qualification"] == qualification and
                    (e["feeStatus"] == fee_status or e["feeStatus"] == "all")), None)
        if est and est.get("derive"):
            dd = self.derive(country, fee_status, est["derive"])
            if dd:
                return dict(dd, status="estimated", note=est.get("basis", ""))
        elif est:
            return {"status": "estimated", "currency": est["currency"], "baseYear": est.get("year") or self.plan,
                    "tuition": num(est.get("tuition")), "otherFees": num(est.get("otherFees")), "admission": num(est.get("admission")),
                    "note": est.get("basis", "")}
        info = self.country(country)
        return {"status": "missing", "currency": info["currency"] if info else "USD", "baseYear": self.plan,
                "tuition": 0, "otherFees": 0, "admission": 0,
                "note": "No published fee found yet for this course in this country. Please enter your own estimate."}

    def bench(self, country, qualification, fee_status, rate):
        return self.verified(country, qualification, fee_status, rate) or self.estimate(country, qualification, fee_status)

    def living_item(self, country, key, fee_status):
        living = next((x for x in self.d.get("living", []) if x["country"] == country), None)
        info = self.country(country)
        missing = {"amount": 0, "currency": info["currency"] if info else "USD", "baseYear": self.plan, "status": "missing"}
        if not living or key not in living.get("items", {}):
            return missing
        it = living["items"][key]
        if it.get("domestic") or it.get("international"):
            it = it.get(fee_status)
            if not it:
                return missing
        return {"amount": num(it.get("amount")), "currency": it.get("currency") or living["currency"],
                "baseYear": it.get("year") or living.get("year") or self.plan, "status": it.get("status") or "verified"}

    def fee_status(self, country, nationality):
        info = self.country(country)
        if not info or not nationality:
            return "international"
        return "domestic" if nationality in (info.get("domesticNationalities") or []) else "international"

    def visa(self, country, nationality):
        v = next((x for x in self.d.get("visas", []) if x["country"] == country), None)
        if self.fee_status(country, nationality) == "domestic":
            return {"amount": 0, "currency": v["currency"] if v else "USD", "baseYear": self.plan, "status": "notNeeded"}
        if not v:
            return {"amount": 0, "currency": "USD", "baseYear": self.plan, "status": "missing"}
        sp = next((x for x in v.get("byNationality", []) if x["nationality"] == nationality), None)
        if sp:
            return {"amount": sp["amount"], "currency": sp.get("currency") or v["currency"], "status": sp.get("status") or "verified",
                    "baseYear": sp.get("year") or v.get("year") or self.plan}
        return {"amount": v["amount"], "currency": v["currency"], "status": v.get("status") or "verified",
                "baseYear": v.get("year") or self.plan}

    def travel(self, country, residence):
        t = self.d.get("travel", {})
        dest = next((x for x in t.get("toCountry", []) if x["country"] == country), None)
        if not dest:
            return {"amount": 0, "currency": "USD", "baseYear": self.plan, "status": "missing"}
        same = bool(residence) and residence == country
        place = next((p for p in self.d.get("places", []) if p["name"] == residence), None)
        region = place["region"] if place else "other"
        fare = dest["domestic"] if same else (dest["byRegion"][region] if isnum(dest["byRegion"].get(region)) else dest["byRegion"]["other"])
        setup = num(t.get("domesticSetup")) if same else num(t.get("setup"))
        return {"amount": fare + setup, "currency": t.get("currency") or "USD", "status": "estimated", "baseYear": t.get("year") or self.plan}


# =====================================================================
# Workbook
# =====================================================================
def build_xlsx(d):
    R = Rules(d)
    wb = Workbook()
    names = {}

    def name(nm, sheet, ref):
        wb.defined_names[nm] = DefinedName(nm, attr_text=f"{q(sheet)}!{ref}")
        names[nm] = True

    S_START, S_IN, S_DASH, S_CMP = "Start Here", "Inputs", "Dashboard", "Compare Countries"
    S_DB, S_REC, S_LIV, S_VT, S_AS, S_SRC, S_LISTS, S_CC = ("Cost Database", "Fee Records", "Living Costs", "Visas and Travel",
                                                          "Assumptions and Currency", "Sources and Method", "Lists", "Compare Calc")
    ws_start = wb.active
    ws_start.title = S_START
    ws_in = wb.create_sheet(S_IN)
    ws_dash = wb.create_sheet(S_DASH)
    ws_cmp = wb.create_sheet(S_CMP)
    kid_ws = [wb.create_sheet("Child %d" % (i + 1)) for i in range(KIDS)]
    ws_db, ws_rec, ws_liv, ws_vt, ws_as, ws_src = (wb.create_sheet(s) for s in (S_DB, S_REC, S_LIV, S_VT, S_AS, S_SRC))
    ws_lists, ws_cc = wb.create_sheet(S_LISTS), wb.create_sheet(S_CC)

    countries = [c["name"] for c in d["countries"]]
    quals = [x["name"] for x in d["qualifications"]]
    places = [p["name"] for p in d["places"]]
    ccys = list(d["currencies"])

    # ---------------- Lists (hidden) ----------------
    ws_lists["A1"], ws_lists["B1"], ws_lists["C1"], ws_lists["D1"], ws_lists["E1"] = "Countries", "Qualifications", "Places", "Currencies", "Domestic (country|nationality)"
    for i, v in enumerate(countries):
        ws_lists.cell(row=2 + i, column=1, value=v)
    for i, v in enumerate(quals):
        ws_lists.cell(row=2 + i, column=2, value=v)
    for i, v in enumerate(places):
        ws_lists.cell(row=2 + i, column=3, value=v)
    for i, v in enumerate(["Automatic"] + ccys):
        ws_lists.cell(row=2 + i, column=4, value=v)
    dom = [c["name"] + "|" + n for c in d["countries"] for n in (c.get("domesticNationalities") or [])]
    for i, v in enumerate(dom):
        ws_lists.cell(row=2 + i, column=5, value=v)
    ws_lists["G1"], ws_lists["H1"], ws_lists["I1"] = "Country", "Currency", "Default living-cost increase"
    for i, c in enumerate(d["countries"]):
        ws_lists.cell(row=2 + i, column=7, value=c["name"])
        ws_lists.cell(row=2 + i, column=8, value=c["currency"])
        ws_lists.cell(row=2 + i, column=9, value=c["livingInflation"])
    nC = len(countries)
    name("DomKey", S_LISTS, f"$E$2:$E${1 + len(dom)}")
    name("Ctry_Name", S_LISTS, f"$G$2:$G${1 + nC}")
    name("Ctry_Ccy", S_LISTS, f"$H$2:$H${1 + nC}")
    name("Ctry_LInf", S_LISTS, f"$I$2:$I${1 + nC}")
    list_countries = f"{q(S_LISTS)}!$A$2:$A${1 + nC}"
    list_quals = f"{q(S_LISTS)}!$B$2:$B${1 + len(quals)}"
    list_places = f"{q(S_LISTS)}!$C$2:$C${1 + len(places)}"
    list_ccys = f"{q(S_LISTS)}!$D$2:$D${2 + len(ccys)}"
    ws_lists.sheet_state = "hidden"

    # ---------------- Assumptions and Currency ----------------
    ws = ws_as
    put(ws, "A1", "Assumptions and currency", TITLE)
    put(ws, "A3", "Plan start year (academic year starting)", BOLD)
    put(ws, "B3", d["planStartYear"])
    name("PlanYear", S_AS, "$B$3")
    put(ws, "A4", "Dataset checked on", BOLD)
    put(ws, "B4", d["datasetDate"])
    put(ws, "A6", "Exchange rates - units of each currency per 1 USD (" + d["exchangeRates"]["date"] + "). Enter your own rate in column C to override.", H2)
    header_row(ws, 7, 1, ["Currency", "Dataset rate", "Your rate (optional)", "Rate used"], [34, 16, 20, 16])
    for i, c in enumerate(ccys):
        r = 8 + i
        put(ws, f"A{r}", c, border=True)
        put(ws, f"B{r}", d["exchangeRates"]["rates"][c], fmt="0.0000", border=True)
        put(ws, f"C{r}", None, fill=INPUT_FILL, fmt="0.0000", lock=False, border=True)
        put(ws, f"D{r}", f"=IF(AND(ISNUMBER(C{r}),C{r}>0),C{r},B{r})", fill=CALC_FILL, fmt="0.0000", border=True)
    last = 7 + len(ccys)
    name("FxCodes", S_AS, f"$A$8:$A${last}")
    name("FxUsed", S_AS, f"$D$8:$D${last}")
    put(ws, f"A{last + 1}", d["exchangeRates"].get("convention", ""), MUTED, wrap=True)
    ws.merge_cells(f"A{last + 1}:D{last + 1}")
    ws.row_dimensions[last + 1].height = 60
    r0 = last + 3
    put(ws, f"A{r0}", "Default investment return by results currency (used when the Inputs return is blank)", H2)
    header_row(ws, r0 + 1, 1, ["Currency", "Yearly return"])
    rets = d["assumptions"]["returnsByCurrency"]
    for i, c in enumerate(rets):
        put(ws, f"A{r0 + 2 + i}", c, border=True)
        put(ws, f"B{r0 + 2 + i}", rets[c], fmt=PCT, border=True)
    name("Ret_Ccy", S_AS, f"$A${r0 + 2}:$A${r0 + 1 + len(rets)}")
    name("Ret_Rate", S_AS, f"$B${r0 + 2}:$B${r0 + 1 + len(rets)}")
    put(ws, f"A{r0 + 2 + len(rets)}", "Any other currency: 5%. " + d["assumptions"].get("returnNote", ""), MUTED, wrap=True)
    ws.merge_cells(f"A{r0 + 2 + len(rets)}:D{r0 + 2 + len(rets)}")
    ws.row_dimensions[r0 + 2 + len(rets)].height = 40

    FX = lambda ccy: f"INDEX(FxUsed,MATCH({ccy},FxCodes,0))"

    # ---------------- Cost Database (one row per country | qualification | fee status) ----------------
    ws = ws_db
    put(ws, "A1", "Cost database - university fees by country, qualification and fee status", TITLE)
    put(ws, "A2", "Published fees are averaged live from the Fee Records sheet. The columns below hold planning estimates "
                  "(used only where no published fee exists), pre-stage fees and course lengths. Read-only; edit "
                  "data/education/education-costs.json and rebuild instead.", MUTED, wrap=True)
    ws.merge_cells("A2:N2")
    ws.row_dimensions[2].height = 30
    cols = ["Key", "Country", "Qualification", "Fee status", "Education currency", "Course length (years)", "Offered (1/0)",
            "Pre-stage years", "Pre-stage qualification", "Default fee increase", "Published records", "Records currency",
            "Estimate status", "Estimate currency", "Estimate base year", "Estimate tuition", "Estimate other fees", "Estimate admission",
            "Pre-stage status", "Pre-stage currency", "Pre-stage base year", "Pre-stage tuition", "Pre-stage other fees", "Pre-stage admission",
            "Basis / note"]
    header_row(ws, 4, 1, cols, [34, 11, 26, 12, 10, 10, 9, 9, 22, 10, 9, 9, 10, 9, 9, 13, 13, 13, 10, 9, 9, 13, 13, 13, 70])
    rec_rows = []
    r = 5
    for c in d["countries"]:
        for qn in quals:
            for fs in ("domestic", "international"):
                key = c["name"] + "|" + qn + "|" + fs
                dur = R.duration(c["name"], qn)
                trate = R.tuition_inflation(c["name"], qn, fs)
                rccy, same = R.used_records(c["name"], qn, fs)
                est = R.estimate(c["name"], qn, fs) if not same else None
                pre = dur["preStage"]
                pb = R.bench(c["name"], pre["qualification"], fs, R.tuition_inflation(c["name"], pre["qualification"], fs)) if pre else None
                row = [key, c["name"], qn, fs, c["currency"], dur["years"] if isnum(dur["years"]) else 0, 1 if dur["available"] else 0,
                       pre["years"] if pre else 0, pre["qualification"] if pre else "", trate, len(same), rccy or "",
                       est["status"] if est else "", est["currency"] if est else "", est["baseYear"] if est else "",
                       est["tuition"] if est else 0, est["otherFees"] if est else 0, est["admission"] if est else 0,
                       pb["status"] if pb else "", pb["currency"] if pb else "", pb["baseYear"] if pb else "",
                       pb["tuition"] if pb else 0, pb["otherFees"] if pb else 0, pb["admission"] if pb else 0,
                       (est["note"] if est else ("Published fees from %d record(s)." % len(same))) + ((" " + dur["note"]) if dur.get("note") else "")]
                for j, v in enumerate(row):
                    cell = ws.cell(row=r, column=1 + j, value=v)
                    cell.font = font(size=9)
                    if j in (5,):
                        cell.number_format = "0.0"
                    if j == 9:
                        cell.number_format = PCT
                    if j in (15, 16, 17, 21, 22, 23):
                        cell.number_format = NUM
                for rec in same:
                    rec_rows.append((key, rec))
                r += 1
    db_last = r - 1
    for nm, col in [("DB_Key", "A"), ("DB_EduCcy", "E"), ("DB_Dur", "F"), ("DB_Avail", "G"), ("DB_PreYears", "H"), ("DB_TInf", "J"),
                    ("DB_RecCcy", "L"), ("DB_EstStatus", "M"), ("DB_EstCcy", "N"), ("DB_EstBase", "O"), ("DB_EstTuition", "P"),
                    ("DB_EstOther", "Q"), ("DB_EstAdm", "R"), ("DB_PreStatus", "S"), ("DB_PreCcy", "T"), ("DB_PreBase", "U"),
                    ("DB_PreTuition", "V"), ("DB_PreOther", "W"), ("DB_PreAdm", "X")]:
        name(nm, S_DB, f"${col}$5:${col}${db_last}")
    ws.freeze_panes = "B5"

    # ---------------- Fee Records (only the records the averages use) ----------------
    ws = ws_rec
    put(ws, "A1", "Published fee records used in the averages", TITLE)
    put(ws, "A2", "Same country, qualification, fee status and currency. Each record is grown to the plan start year with the "
                  "fee increase before averaging.", MUTED)
    header_row(ws, 4, 1, ["Key", "University", "Fee year", "Fee year start", "Currency", "Tuition (per year)",
                          "Other mandatory fees (per year)", "One-time fees", "Status", "Source"],
               [34, 34, 10, 9, 9, 14, 14, 12, 12, 60])
    for i, (key, rec) in enumerate(rec_rows):
        rr = 5 + i
        vals = [key, rec.get("university", ""), rec.get("feeYear", ""), rec.get("feeYearStart") or d["planStartYear"], rec["currency"],
                rec["tuition"], num(rec.get("otherMandatoryAnnual")), num(rec.get("oneTimeAdmission")) + num(rec.get("otherOneTime")),
                rec.get("status", ""), rec.get("sourceUrl", "")]
        for j, v in enumerate(vals):
            cell = ws.cell(row=rr, column=1 + j, value=v)
            cell.font = font(size=9)
            if j in (5, 6, 7):
                cell.number_format = NUM
    rec_last = 4 + len(rec_rows)
    for nm, col in [("RecKey", "A"), ("RecYear", "D"), ("RecTuition", "F"), ("RecOther", "G"), ("RecAdm", "H")]:
        name(nm, S_REC, f"${col}$5:${col}${rec_last}")
    ws.freeze_panes = "B5"

    # ---------------- Living Costs (country | fee status) ----------------
    ws = ws_liv
    put(ws, "A1", "Living costs per year (today's prices)", TITLE)
    put(ws, "A2", d.get("livingNote", ""), MUTED, wrap=True)
    ws.merge_cells("A2:L2")
    ws.row_dimensions[2].height = 40
    hdr = ["Key (country|fee status)"]
    for k in LIVING_KEYS:
        lab = next(it[1] for it in ITEMS if it[0] == k)
        hdr += [lab, lab + " currency", lab + " year", lab + " status"]
    header_row(ws, 4, 1, hdr, [24] + [12, 8, 7, 10] * len(LIVING_KEYS))
    r = 5
    for c in d["countries"]:
        for fs in ("domestic", "international"):
            ws.cell(row=r, column=1, value=c["name"] + "|" + fs).font = font(size=9)
            for j, k in enumerate(LIVING_KEYS):
                it = R.living_item(c["name"], k, fs)
                for m, v in enumerate([it["amount"], it["currency"], it["baseYear"], it["status"]]):
                    cell = ws.cell(row=r, column=2 + j * 4 + m, value=v)
                    cell.font = font(size=9)
                    if m == 0:
                        cell.number_format = NUM
            r += 1
    liv_last = r - 1
    name("LivKey", S_LIV, f"$A$5:$A${liv_last}")
    for j, k in enumerate(LIVING_KEYS):
        for m, part in enumerate(["Amt", "Ccy", "Base", "Status"]):
            name(f"Liv_{k}_{part}", S_LIV, f"${L(2 + j * 4 + m)}$5:${L(2 + j * 4 + m)}${liv_last}")

    # ---------------- Visas and Travel ----------------
    ws = ws_vt
    put(ws, "A1", "Student visa (by destination and nationality) and first travel (by destination and country of residence)", TITLE)
    header_row(ws, 3, 1, ["Key (country|nationality)", "Visa amount", "Currency", "Year", "Status"], [30, 12, 9, 7, 11])
    header_row(ws, 3, 7, ["Key (country|residence)", "Travel and setting-up", "Currency", "Year", "Status"], [30, 12, 9, 7, 11])
    r = 4
    for c in countries:
        for p in places:
            v, t = R.visa(c, p), R.travel(c, p)
            for j, val in enumerate([c + "|" + p, v["amount"], v["currency"], v["baseYear"], v["status"]]):
                ws.cell(row=r, column=1 + j, value=val).font = font(size=9)
            for j, val in enumerate([c + "|" + p, t["amount"], t["currency"], t["baseYear"], t["status"]]):
                ws.cell(row=r, column=7 + j, value=val).font = font(size=9)
            r += 1
    vt_last = r - 1
    for nm, col in [("VisaKey", "A"), ("VisaAmt", "B"), ("VisaCcy", "C"), ("VisaBase", "D"), ("VisaStatus", "E"),
                    ("TravKey", "G"), ("TravAmt", "H"), ("TravCcy", "I"), ("TravBase", "J"), ("TravStatus", "K")]:
        name(nm, S_VT, f"${col}$4:${col}${vt_last}")
    put(ws, "M3", "Visa notes", BOLD)
    for i, v in enumerate(d.get("visas", [])):
        put(ws, f"M{4 + i}", v["country"] + ": " + v.get("note", ""), font(size=9), wrap=True)
    put(ws, f"M{5 + len(d.get('visas', []))}", "Travel: " + d.get("travel", {}).get("note", ""), font(size=9), wrap=True)
    ws.column_dimensions["M"].width = 90

    # ---------------- Inputs ----------------
    ws = ws_in
    ws.column_dimensions["A"].width = 2
    ws.column_dimensions["B"].width = 64
    for c in "CDEF":
        ws.column_dimensions[c].width = 22
    ws.column_dimensions["G"].width = 60
    put(ws, "B1", "Children's Future Education Fund - your inputs", TITLE)
    put(ws, "B2", "Fill in the yellow cells. Leave a cell blank to use the researched figure or default. Results are on the Dashboard.", MUTED)
    fam = [("Number of children (1-4)", 1, "0", "Children beyond this number are ignored."),
           ("Country where your family lives", "Pakistan", None, "Used for first-travel estimates."),
           ("Children's nationality", "Pakistan", None, "Decides local or international fees and visa costs."),
           ("Show results in (currency)", "USD", None, "'Automatic' uses Child 1's study country currency."),
           ("Covered by scholarship / part-time work", 0, PCT, "Reduces tuition, university fees and living costs. Visa and travel are not reduced."),
           ("Investment return per year (blank = default for the currency)", None, PCT, "An assumption, not a guarantee."),
           ("Yearly increase in your savings", 0, PCT, "0% = the same amount every year."),
           ("Compare countries for child number", 1, "0", "Used on the Compare Countries sheet.")]
    put(ws, "B3", "Family", H2)
    for i, (lab, val, fmt, note) in enumerate(fam):
        r = 4 + i
        put(ws, f"B{r}", lab, BOLD)
        put(ws, f"C{r}", val, fill=INPUT_FILL, fmt=fmt, lock=False, border=True)
        put(ws, f"G{r}", note, MUTED)
    put(ws, "D3", "Value used", MUTED)
    put(ws, "D7", '=IF(C7="Automatic",IFERROR(INDEX(Ctry_Ccy,MATCH(C17,Ctry_Name,0)),"USD"),IF(C7="","USD",C7))', fill=CALC_FILL)
    put(ws, "D8", "=MIN(1,MAX(0,N(C8)))", fill=CALC_FILL, fmt=PCT)
    put(ws, "D9", "=IF(ISNUMBER(C9),C9,IFERROR(INDEX(Ret_Rate,MATCH(D7,Ret_Ccy,0)),0.05))", fill=CALC_FILL, fmt=PCT)
    put(ws, "D10", "=N(C10)", fill=CALC_FILL, fmt=PCT)
    put(ws, "D4", "=MIN(4,MAX(1,ROUND(N(C4),0)))", fill=CALC_FILL, fmt="0")
    put(ws, "D11", "=MIN(D4,MAX(1,ROUND(N(C11),0)))", fill=CALC_FILL, fmt="0")
    name("NumChildren", S_IN, "$D$4")
    name("Residence", S_IN, "$C$5")
    name("Nationality", S_IN, "$C$6")
    name("RepCcy", S_IN, "$D$7")
    name("Coverage", S_IN, "$D$8")
    name("ReturnRate", S_IN, "$D$9")
    name("SavInc", S_IN, "$D$10")
    name("CmpChild", S_IN, "$D$11")

    put(ws, "B13", "Children", H2)
    for i in range(KIDS):
        put(ws, f"{L(3 + i)}13", "Child %d" % (i + 1), HDR, fill=HDR_FILL, border=True)
    kid_rows = [("name", "Name (optional)", None, None), ("age", "Current age (0-18)", None, "0"),
                ("schoolClass", "Current school class (0-13, for your reference)", None, "0"),
                ("country", "Education country", None, None), ("qualification", "Qualification", None, None),
                ("entryAge", "College-entry age (usually 18)", 18, "0"),
                ("tInf", "Your own fee increase per year (blank = researched default)", None, PCT),
                ("lInf", "Your own living-cost increase per year (blank = default)", None, PCT)]
    IN = {}
    for j, (key, lab, default, fmt) in enumerate(kid_rows):
        r = 14 + j
        put(ws, f"B{r}", lab, BOLD)
        IN[key] = r
        for i in range(KIDS):
            put(ws, f"{L(3 + i)}{r}", default, fill=INPUT_FILL, fmt=fmt, lock=False, border=True)
    put(ws, "B22", "Your own cost figures (optional) - per year unless one-time, today's prices, in the education country's currency", H2)
    ws.merge_cells("B22:F22")
    ws.row_dimensions[22].height = 34
    ws["B22"].alignment = WRAP
    for j, (key, lab) in enumerate([(it[0], it[1] + (" (one-time)" if it[3] == "once" else "")) for it in ITEMS] + PRE_ITEMS):
        r = 23 + j
        put(ws, f"B{r}", lab)
        IN["ov_" + key] = r
        for i in range(KIDS):
            put(ws, f"{L(3 + i)}{r}", None, fill=INPUT_FILL, fmt=NUM, lock=False, border=True)
    put(ws, "B36", "Researched figures and their sources are on each child's sheet (Child 1 to Child 4). Figures marked "
                   "'Estimated — please review' are planning estimates, not official fees.", MUTED, wrap=True)
    ws.merge_cells("B36:G36")
    ws.row_dimensions[36].height = 30

    def dv(formula, ref, kind="list", **kw):
        v = DataValidation(type=kind, formula1=formula, allow_blank=True, **kw)
        v.error, v.errorTitle = "Please choose or enter a valid value.", "Invalid value"
        ws.add_data_validation(v)
        v.add(ref)
    dv('"1,2,3,4"', "C4")
    dv("=" + list_places, "C5")
    dv("=" + list_places, "C6")
    dv("=" + list_ccys, "C7")
    dv("0", "C8", "decimal", operator="between", formula2="1")
    dv("-0.5", "C9", "decimal", operator="between", formula2="0.3")
    dv("0", "C10", "decimal", operator="between", formula2="0.2")
    dv('"1,2,3,4"', "C11")
    dv("0", f"C{IN['age']}:F{IN['age']}", "whole", operator="between", formula2="18")
    dv("0", f"C{IN['schoolClass']}:F{IN['schoolClass']}", "whole", operator="between", formula2="13")
    dv("=" + list_countries, f"C{IN['country']}:F{IN['country']}")
    dv("=" + list_quals, f"C{IN['qualification']}:F{IN['qualification']}")
    dv("14", f"C{IN['entryAge']}:F{IN['entryAge']}", "whole", operator="between", formula2="45")
    dv("-0.2", f"C{IN['tInf']}:F{IN['lInf']}", "decimal", operator="between", formula2="0.5")
    dv("0", f"C23:F34", "decimal", operator="greaterThanOrEqual")
    ws.freeze_panes = "C4"

    # ---------------- cost block (used by each child sheet and by Compare Calc) ----------------
    def cost_block(ws, top, p):
        """Writes one child's cost calculation from row `top`. p: formulas for active, country, qual, age, entry,
        tinf, linf (custom rate cells or None) and ov (dict item -> cell, or None). Returns cell addresses."""
        A = {}
        rows = ["Included (1 = yes)", "Education country", "Qualification", "Current age", "College-entry age", "Fee status",
                "Lookup key", "Cost database row", "Education currency", "Course length (years)", "Offered in this country (1/0)",
                "Pre-stage years", "Fee increase used", "Living-cost increase used", "Years until college", "Study years",
                "Plan years", "Saving years", "Calculated (1 = yes)", "Message", "Exchange rate (education to results currency)",
                "Published fee records used", "Fee records currency", "Living costs row", "Visa row", "Travel row"]
        for i, lab in enumerate(rows):
            A[lab] = f"$B${top + i}"
            put(ws, f"A{top + i}", lab)
        a = A
        b = lambda lab: a[lab]
        age, entry = b("Current age"), b("College-entry age")
        age_ok = f"IFERROR(AND(ISNUMBER({age}),{age}>=0,{age}<=18,INT({age})={age}),FALSE)"
        entry_ok = f"IFERROR(AND(ISNUMBER({entry}),{entry}>=14,{entry}<=45,INT({entry})={entry}),FALSE)"
        F_ = {
            "Included (1 = yes)": "=" + p["active"],
            "Education country": f'=IF({p["country"]}="","",{p["country"]})',
            "Qualification": f'=IF({p["qual"]}="","",{p["qual"]})',
            "Current age": f'=IF({p["age"]}="","",{p["age"]})',
            "College-entry age": f'=IF({p["entry"]}="","",{p["entry"]})',
            "Fee status": f'=IF({b("Education country")}="","",IF(COUNTIF(DomKey,{b("Education country")}&"|"&Nationality)>0,"domestic","international"))',
            "Lookup key": f'={b("Education country")}&"|"&{b("Qualification")}&"|"&{b("Fee status")}',
            "Cost database row": f'=IFERROR(MATCH({b("Lookup key")},DB_Key,0),0)',
            "Education currency": f'=IFERROR(INDEX(Ctry_Ccy,MATCH({b("Education country")},Ctry_Name,0)),"USD")',
            "Course length (years)": f'=IF({b("Cost database row")}>0,INDEX(DB_Dur,{b("Cost database row")}),0)',
            "Offered in this country (1/0)": f'=IF({b("Cost database row")}>0,INDEX(DB_Avail,{b("Cost database row")}),0)',
            "Pre-stage years": f'=IF({b("Cost database row")}>0,INDEX(DB_PreYears,{b("Cost database row")}),0)',
            "Fee increase used": (f'=IF(ISNUMBER({p["tinf"]}),{p["tinf"]},' if p["tinf"] else "=IF(FALSE,0,") +
                                 f'IF({b("Cost database row")}>0,INDEX(DB_TInf,{b("Cost database row")}),0))',
            "Living-cost increase used": (f'=IF(ISNUMBER({p["linf"]}),{p["linf"]},' if p["linf"] else "=IF(FALSE,0,") +
                                         f'IFERROR(INDEX(Ctry_LInf,MATCH({b("Education country")},Ctry_Name,0)),0))',
            "Years until college": f'=IF({b("Calculated (1 = yes)")}=1,MAX(0,{entry}-{age}),0)',
            "Study years": f'=ROUNDUP({b("Course length (years)")},0)',
            "Plan years": f'=IF({b("Calculated (1 = yes)")}=1,{b("Years until college")}+{b("Study years")},0)',
            "Saving years": f'=MAX(0,{b("Plan years")}-1)',
            "Calculated (1 = yes)": f'=IF(IFERROR(AND({b("Included (1 = yes)")}=1,{b("Education country")}<>"",{b("Qualification")}<>"",'
                                    f'{b("Offered in this country (1/0)")}=1,{age_ok},{entry_ok}),FALSE),1,0)',
            "Message": f'=IF({b("Included (1 = yes)")}<>1,"",IF(AND({age}="",{b("Education country")}="",{b("Qualification")}=""),"Enter this child\'s details on the Inputs sheet.",IF(NOT({age_ok}),"Choose the child\'s current age (0 to 18).",'
                       f'IF(NOT({entry_ok}),"College-entry age must be a whole number from 14 to 45.",'
                       f'IF({b("Education country")}="","Choose the education country.",IF({b("Qualification")}="","Choose the qualification.",'
                       f'IF({b("Offered in this country (1/0)")}<>1,"This qualification is not offered (or has no typical duration) in the selected country.","")))))))',
            "Exchange rate (education to results currency)": f'={FX("RepCcy")}/{FX(b("Education currency"))}',
            "Published fee records used": f'=IF({b("Cost database row")}=0,0,COUNTIF(RecKey,{b("Lookup key")}))',
            "Fee records currency": f'=IF({b("Published fee records used")}>0,INDEX(DB_RecCcy,{b("Cost database row")}),'
                                    f'IF({b("Cost database row")}>0,INDEX(DB_EstCcy,{b("Cost database row")}),{b("Education currency")}))',
            "Living costs row": f'=IFERROR(MATCH({b("Education country")}&"|"&{b("Fee status")},LivKey,0),0)',
            "Visa row": f'=IFERROR(MATCH({b("Education country")}&"|"&Nationality,VisaKey,0),0)',
            "Travel row": f'=IFERROR(MATCH({b("Education country")}&"|"&Residence,TravKey,0),0)',
        }
        for lab in rows:
            cell = put(ws, a[lab].replace("$", ""), F_[lab], fill=CALC_FILL)
            if lab in ("Fee increase used", "Living-cost increase used"):
                cell.number_format = PCT
        row_db, cnt, rccy, edu = b("Cost database row"), b("Published fee records used"), b("Fee records currency"), b("Education currency")
        trate, pre_y, lrow, vrow, trow = b("Fee increase used"), b("Pre-stage years"), b("Living costs row"), b("Visa row"), b("Travel row")
        conv = lambda amt, ccy: f"({amt})/{FX(ccy)}*{FX(edu)}"

        # cost lines (today's prices, education currency)
        ct = top + len(rows) + 1
        header_row(ws, ct, 1, ["Cost line", "Researched (education currency)", "Your figure", "Amount used", "Price year",
                               "Status code", "Status", "Increase", "Timing", "Scholarship applies"])
        lines = {}
        allitems = [(k, lab) for k, lab, *_ in ITEMS] + PRE_ITEMS
        for i, (k, lab) in enumerate(allitems):
            r = ct + 1 + i
            lines[k] = r
            put(ws, f"A{r}", lab)
            def main_bench(rec_col, est_col):
                return (f"IF({row_db}=0,0,{conv(f'IF({cnt}>0,SUMPRODUCT((RecKey={b('Lookup key')})*{rec_col}*(1+{trate})^(PlanYear-RecYear))/{cnt},INDEX({est_col},{row_db}))', rccy)})")
            main_base = f"IF({cnt}>0,PlanYear,IF({row_db}>0,INDEX(DB_EstBase,{row_db}),PlanYear))"
            main_status = f'IF({row_db}=0,"missing",IF({cnt}>0,"verified",INDEX(DB_EstStatus,{row_db})))'
            pre_amt = lambda col: f"IF({pre_y}>0,{conv(f'INDEX({col},{row_db})', f'INDEX(DB_PreCcy,{row_db})')},0)"
            if k == "tuition":
                res, base, st = "=" + main_bench("RecTuition", "DB_EstTuition"), "=" + main_base, "=" + main_status
            elif k == "otherFees":
                res, base, st = "=" + main_bench("RecOther", "DB_EstOther"), "=" + main_base, "=" + main_status
            elif k == "admission":
                res = f"=IF({pre_y}>0,{pre_amt('DB_PreAdm')},{main_bench('RecAdm', 'DB_EstAdm')})"
                base = f"=IF({pre_y}>0,INDEX(DB_PreBase,{row_db}),{main_base})"
                st = f"=IF({pre_y}>0,INDEX(DB_PreStatus,{row_db}),{main_status})"
            elif k in ("preTuition", "preOtherFees"):
                col = "DB_PreTuition" if k == "preTuition" else "DB_PreOther"
                res = "=" + pre_amt(col)
                base = f"=IF({pre_y}>0,INDEX(DB_PreBase,{row_db}),PlanYear)"
                st = f'=IF({pre_y}>0,INDEX(DB_PreStatus,{row_db}),"")'
            elif k in LIVING_KEYS:
                res = f"=IF({lrow}=0,0,{conv(f'INDEX(Liv_{k}_Amt,{lrow})', f'INDEX(Liv_{k}_Ccy,{lrow})')})"
                base = f"=IF({lrow}=0,PlanYear,INDEX(Liv_{k}_Base,{lrow}))"
                st = f'=IF({lrow}=0,"missing",INDEX(Liv_{k}_Status,{lrow}))'
            elif k == "visaApplication":
                res = f"=IF({vrow}=0,0,{conv(f'INDEX(VisaAmt,{vrow})', f'INDEX(VisaCcy,{vrow})')})"
                base = f"=IF({vrow}=0,PlanYear,INDEX(VisaBase,{vrow}))"
                st = f'=IF({vrow}=0,"missing",INDEX(VisaStatus,{vrow}))'
            else:  # travelRelocation
                res = f"=IF({trow}=0,0,{conv(f'INDEX(TravAmt,{trow})', f'INDEX(TravCcy,{trow})')})"
                base = f"=IF({trow}=0,PlanYear,INDEX(TravBase,{trow}))"
                st = f'=IF({trow}=0,"missing",INDEX(TravStatus,{trow}))'
            ov = p["ov"][k] if p["ov"] else None
            put(ws, f"B{r}", res, fill=CALC_FILL, fmt=NUM)
            put(ws, f"C{r}", ("=IF(ISNUMBER(%s),%s,\"\")" % (ov, ov)) if ov else "", fmt=NUM)
            put(ws, f"D{r}", f"=IF(ISNUMBER(C{r}),IF(C{r}>=0,C{r},B{r}),B{r})", fill=CALC_FILL, fmt=NUM)
            put(ws, f"E{r}", f"=IF(ISNUMBER(C{r}),PlanYear,{base[1:]})", fill=CALC_FILL)
            put(ws, f"F{r}", f'=IF(ISNUMBER(C{r}),"override",{st[1:]})', fill=CALC_FILL)
            put(ws, f"G{r}", "=" + STATUS_LABEL.format(s=f"F{r}"))
            meta = next((it for it in ITEMS if it[0] == k), None)
            put(ws, f"H{r}", meta[2] if meta else "tuition")
            put(ws, f"I{r}", meta[3] if meta else "annual")
            put(ws, f"J{r}", ("Yes" if meta[4] else "No") if meta else "Yes")

        # study years: gross and net (after scholarship) per cost line
        st_top = ct + len(allitems) + 2
        hdr = ["Study year (0 = first)", "Calendar year", "Fraction of year", "In pre-stage (1/0)", "In course (1/0)"]
        hdr += [lab + " (gross)" for _, lab, *_ in ITEMS] + [lab for _, lab, *_ in ITEMS] + ["Education expenses", "Before scholarship"]
        header_row(ws, st_top, 1, hdr)
        calc_ok, n_ref, E_ref, D_ref = b("Calculated (1 = yes)"), b("Years until college"), b("Study years"), b("Course length (years)")
        fx = b("Exchange rate (education to results currency)")
        lrate = b("Living-cost increase used")
        first = st_top + 1
        for kk in range(STUDY):
            r = first + kk
            put(ws, f"A{r}", kk)
            put(ws, f"B{r}", f"=PlanYear+{n_ref}+A{r}")
            put(ws, f"C{r}", f"=MAX(0,MIN(1,{D_ref}-A{r}))")
            put(ws, f"D{r}", f"=IF(A{r}<{pre_y},1,0)")
            put(ws, f"E{r}", f"=IF(AND({calc_ok}=1,A{r}<{E_ref}),1,0)")
            for j, (k, lab, inf, timing, covered) in enumerate(ITEMS):
                if k == "tuition":
                    amt, base = f"IF(D{r}=1,$D${lines['preTuition']},$D${lines['tuition']})", f"IF(D{r}=1,$E${lines['preTuition']},$E${lines['tuition']})"
                elif k == "otherFees":
                    amt, base = f"IF(D{r}=1,$D${lines['preOtherFees']},$D${lines['otherFees']})", f"IF(D{r}=1,$E${lines['preOtherFees']},$E${lines['otherFees']})"
                else:
                    amt, base = f"$D${lines[k]}", f"$E${lines[k]}"
                rate = trate if inf == "tuition" else lrate
                when = f"E{r}=1" if timing == "annual" else f"E{r}=1,A{r}=0"
                frac = f"*C{r}" if timing == "annual" else ""
                g = L(6 + j)
                put(ws, f"{g}{r}", f"=IF(AND({when}),{amt}*(1+{rate})^(B{r}-{base}){frac}*{fx},0)", fmt=NUM)
                put(ws, f"{L(6 + len(ITEMS) + j)}{r}", f"={g}{r}*{'(1-Coverage)' if covered else '1'}", fmt=NUM)
            net0, net1 = L(6 + len(ITEMS)), L(5 + 2 * len(ITEMS))
            put(ws, f"{L(6 + 2 * len(ITEMS))}{r}", f"=SUM({net0}{r}:{net1}{r})", BOLD, fmt=NUM)
            put(ws, f"{L(7 + 2 * len(ITEMS))}{r}", f"=SUM(F{r}:{L(5 + len(ITEMS))}{r})", fmt=NUM)
        last_r = first + STUDY - 1
        tot = last_r + 1
        put(ws, f"A{tot}", "Total", BOLD)
        for j in range(2 * len(ITEMS) + 2):
            col = L(6 + j)
            put(ws, f"{col}{tot}", f"=SUM({col}{first}:{col}{last_r})", BOLD, fmt=NUM)
        exp_col = L(6 + 2 * len(ITEMS))
        out = {
            "ok": calc_ok, "n": n_ref, "E": E_ref, "H": b("Plan years"), "P": b("Saving years"), "D": D_ref, "msg": b("Message"),
            "country": b("Education country"), "qual": b("Qualification"),
            "exp_range": f"${exp_col}${first}:${exp_col}${last_r}",
            "total": f"${exp_col}${tot}", "gross": f"${L(7 + 2 * len(ITEMS))}${tot}",
            "item_total": {k: f"${L(6 + len(ITEMS) + j)}${tot}" for j, (k, *_r) in enumerate(ITEMS)},
            "status": {k: f"$F${lines[k]}" for k in lines}, "pre_y": pre_y, "end": tot,
        }
        # summary
        s = tot + 2
        put(ws, f"A{s}", "Total estimated education cost (results currency)", BOLD)
        put(ws, f"B{s}", f"={out['total']}", BOLD, fill=GOLD_FILL, fmt=NUM)
        put(ws, f"A{s + 1}", "Covered by scholarship / part-time work")
        put(ws, f"B{s + 1}", f"={out['gross']}-{out['total']}", fmt=NUM)
        put(ws, f"A{s + 2}", "Tuition figure status")
        ts, ps = out["status"]["tuition"], out["status"]["preTuition"]
        put(ws, f"B{s + 2}", f'=IF(OR({ts}="missing",AND({pre_y}>0,{ps}="missing")),"missing",IF(OR({ts}="estimated",AND({pre_y}>0,{ps}="estimated")),"estimated",IF({pre_y}>0,{ps},{ts})))')
        put(ws, f"A{s + 3}", "Some costs have no figure (1/0)")
        put(ws, f"B{s + 3}", f'=IF(OR({ts}="missing",AND({pre_y}>0,{ps}="missing"),{out["status"]["accommodation"]}="missing"),1,0)')
        out.update({"tuition_status": f"$B${s + 2}", "has_missing": f"$B${s + 3}", "covered": f"$B${s + 1}", "end": s + 3})
        return out

    # ---------------- Child sheets: cost block + savings + yearly roll-forward ----------------
    KO = []
    for i, ws in enumerate(kid_ws):
        col = L(3 + i)
        ref = lambda key: f"{q(S_IN)}!${col}${IN[key]}"
        ws.column_dimensions["A"].width = 44
        ws.column_dimensions["B"].width = 18
        for c in range(3, 30):
            ws.column_dimensions[L(c)].width = 13
        ws.column_dimensions["G"].width = 30
        put(ws, "A1", "Child %d - cost and savings calculation" % (i + 1), TITLE)
        put(ws, "A2", "All formulas. Change inputs on the Inputs sheet. Amounts in the cost lines table are in the education "
                      "country's currency; everything else is in the results currency.", MUTED)
        o = cost_block(ws, 4, {"active": f"IF({i + 1}<=NumChildren,1,0)", "country": ref("country"), "qual": ref("qualification"),
                                "age": ref("age"), "entry": ref("entryAge"), "tinf": ref("tInf"), "linf": ref("lInf"),
                                "ov": {k: ref("ov_" + k) for k in [it[0] for it in ITEMS] + [x[0] for x in PRE_ITEMS]}})
        # savings and roll-forward
        top = o["end"] + 3
        put(ws, f"A{top - 1}", "Savings and yearly fund (results currency). Costs are paid at the start of each year, savings added at the end.", H2)
        hdr = ["Plan year", "Calendar year", "Study year", "Education expenses", "Saving factor", "Growth factor",
               "a (no saving)", "b (per 1 saved)", "A carried", "B carried", "Needed-now candidate", "Saving candidate",
               "Opening fund", "Family savings added", "Investment growth", "Closing fund", "Studying (1/0)"]
        header_row(ws, top, 1, hdr)
        first = top + 1
        lastp = first + ROWS - 1
        lump_c, x_c = f"$B${lastp + 2}", f"$B${lastp + 3}"
        ok, n_, E_, H_, P_ = o["ok"], o["n"], o["E"], o["H"], o["P"]
        for pp in range(1, ROWS + 1):
            r = first + pp - 1
            prev = r - 1
            put(ws, f"A{r}", pp)
            put(ws, f"B{r}", f"=PlanYear+A{r}-1")
            put(ws, f"C{r}", f"=A{r}-1-{n_}")
            put(ws, f"D{r}", f"=IF(AND({ok}=1,C{r}>=0,C{r}<{E_}),INDEX({o['exp_range']},C{r}+1),0)", fmt=NUM)
            put(ws, f"E{r}", f"=IF(AND({ok}=1,A{r}<={P_}),(1+SavInc)^(A{r}-1),0)", fmt="0.0000")
            put(ws, f"F{r}", f"=(1+ReturnRate)^(A{r}-1)", fmt="0.0000")
            put(ws, f"G{r}", (f"=I{prev}-D{r}" if pp > 1 else f"=-D{r}"), fmt=NUM)
            put(ws, f"H{r}", (f"=J{prev}" if pp > 1 else "=0"), fmt="0.0000")
            put(ws, f"I{r}", f"=G{r}*(1+ReturnRate)", fmt=NUM)
            put(ws, f"J{r}", f"=H{r}*(1+ReturnRate)+E{r}", fmt="0.0000")
            put(ws, f"K{r}", f"=IF(AND(D{r}>0,H{r}<=1E-12,G{r}<0),-G{r}/F{r},0)", fmt=NUM)
            put(ws, f"L{r}", f"=IF(AND(D{r}>0,H{r}>1E-12),-(G{r}+{lump_c}*F{r})/H{r},0)", fmt=NUM)
            put(ws, f"M{r}", (f"=P{prev}" if pp > 1 else f"={lump_c}"), fmt=NUM)
            put(ws, f"N{r}", f"=E{r}*{x_c}", fmt=NUM)
            put(ws, f"O{r}", f"=IF(AND({ok}=1,A{r}<={H_}),(M{r}-D{r})*ReturnRate,0)", fmt=NUM)
            put(ws, f"P{r}", f"=M{r}-D{r}+O{r}+N{r}", fmt=NUM)
            put(ws, f"Q{r}", f"=IF(AND({ok}=1,C{r}>=0,C{r}<{E_}),1,0)")
        put(ws, f"A{lastp + 2}", "Amount needed now", BOLD)
        put(ws, lump_c.replace("$", ""), f"=IF({ok}=1,MAX(K{first}:K{lastp}),0)", BOLD, fill=GOLD_FILL, fmt=NUM)
        put(ws, f"A{lastp + 3}", "Required yearly saving (first year)", BOLD)
        put(ws, x_c.replace("$", ""), f"=IF({ok}=1,IF(MAX(0,MAX(L{first}:L{lastp}))<1E-9,0,MAX(0,MAX(L{first}:L{lastp}))),0)",
            BOLD, fill=GOLD_FILL, fmt=NUM)
        ws.freeze_panes = "B4"
        o.update({"sheet": ws.title, "first": first, "lump": lump_c, "x": x_c})
        KO.append(o)

    def kref(i, cell):
        return f"{q(KO[i]['sheet'])}!{cell}"

    def kcol(i, col, pp):
        return f"{q(KO[i]['sheet'])}!${col}${KO[i]['first'] + pp - 1}"

    kname = lambda i: f'IF({q(S_IN)}!${L(3 + i)}${IN["name"]}<>"",{q(S_IN)}!${L(3 + i)}${IN["name"]},"Child {i + 1}")'

    # ---------------- Dashboard ----------------
    ws = ws_dash
    ws.column_dimensions["A"].width = 2
    widths = {"B": 44, "C": 16, "D": 26, "E": 18, "F": 18, "G": 20, "H": 18, "I": 50}
    for c, w in widths.items():
        ws.column_dimensions[c].width = w
    put(ws, "B1", "Children's Future Education Fund - dashboard", TITLE)
    put(ws, "B2", '="All amounts in "&RepCcy&", at future prices. Planning estimates only - not financial advice."', MUTED)
    total_c, save_c, lump_c, hmax_c = "$C$4", "$C$5", "$C$6", "$C$8"
    put(ws, "B4", "Total education fund required", BIG)
    put(ws, "C4", "=" + "+".join(kref(i, KO[i]["total"]) for i in range(KIDS)), BIG, fill=GOLD_FILL, fmt=NUM)
    put(ws, "B5", "Required yearly savings (this year)", BIG)
    put(ws, "C5", "=" + "+".join(kcol(i, "N", 1) for i in range(KIDS)), BIG, fill=GOLD_FILL, fmt=NUM)
    put(ws, "B6", "Amount needed now", BOLD)
    put(ws, "C6", "=" + "+".join(kref(i, KO[i]["lump"]) for i in range(KIDS)), BOLD, fmt=NUM)
    put(ws, "D4", '=IF(C4=0,"Enter your family and children\'s details on the Inputs sheet to see the figures.","")', MUTED)
    put(ws, "D6", '=IF(C6>0.5,"Plus about "&TEXT(C6,"#,##0")&" "&RepCcy&" needed now, because a course starts before any yearly saving can be made.","")', MUTED)
    put(ws, "B7", "Covered by scholarship / part-time work")
    put(ws, "C7", "=" + "+".join(kref(i, KO[i]["covered"]) for i in range(KIDS)), fmt=NUM)
    put(ws, "B8", "Plan length (years)")
    put(ws, "C8", "=MAX(" + ",".join(kref(i, KO[i]["H"]) for i in range(KIDS)) + ")")

    put(ws, "B10", "Each child", H2)
    header_row(ws, 11, 2, ["Child", "Education country", "Qualification", "Course length (years)", "Education starts in",
                           '="Total ("&RepCcy&")"', "Message"])
    for i in range(KIDS):
        r = 12 + i
        on = f"{i + 1}<=NumChildren"
        o = KO[i]
        put(ws, f"B{r}", f'=IF({on},{kname(i)},"")')
        put(ws, f"C{r}", f'=IF({on},{kref(i, o["country"])},"")')
        put(ws, f"D{r}", f'=IF({on},{kref(i, o["qual"])},"")')
        put(ws, f"E{r}", f'=IF(AND({on},{kref(i, o["ok"])}=1),{kref(i, o["D"])},"")', fmt="0.0")
        put(ws, f"F{r}", f'=IF(AND({on},{kref(i, o["ok"])}=1),IF({kref(i, o["n"])}=0,"This year",{kref(i, o["n"])}&" years")&" ("&(PlanYear+{kref(i, o["n"])})&"/"&RIGHT(PlanYear+{kref(i, o["n"])}+1,2)&")","")')
        put(ws, f"G{r}", f'=IF({on},{kref(i, o["total"])},"")', BOLD, fmt=NUM)
        put(ws, f"H{r}", f'=IF({on},{kref(i, o["msg"])},"")', MUTED)
    put(ws, "B16", "All children", BOLD)
    put(ws, "G16", "=C4", BOLD, fmt=NUM)

    put(ws, "B18", "Where the money goes", H2)
    header_row(ws, 19, 2, ["Cost type", '="Amount ("&RepCcy&")"', "Share"])
    for j, (lab, keys) in enumerate(GROUPS):
        r = 20 + j
        put(ws, f"B{r}", lab)
        put(ws, f"C{r}", "=" + "+".join(kref(i, KO[i]["item_total"][k]) for i in range(KIDS) for k in keys), fmt=NUM)
        put(ws, f"D{r}", f"=IF($C$4>0,C{r}/$C$4,0)", fmt="0%")
    put(ws, "B28", "Total", BOLD)
    put(ws, "C28", "=SUM(C20:C27)", BOLD, fmt=NUM)

    put(ws, "B30", "Year-by-year plan - how the education fund builds up and is spent", H2)
    yt = 31
    header_row(ws, yt, 2, ["Year", "Studying", "Opening fund", "Family savings added", "Investment growth", "Education expenses", "Closing fund"])
    for pp in range(1, ROWS + 1):
        r = yt + pp
        show = f"{pp}<={hmax_c}"
        put(ws, f"B{r}", f'=IF({show},(PlanYear+{pp - 1})&"/"&RIGHT(PlanYear+{pp},2),"")')
        names_ = "&".join(f'IF({kcol(i, "Q", pp)}=1,", "&{kname(i)},"")' for i in range(KIDS))
        put(ws, f"C{r}", f'=IF({show},MID({names_},3,200),"")')
        for j, colk in enumerate(["M", "N", "O", "D", "P"]):
            put(ws, f"{L(4 + j)}{r}", f'=IF({show},ROUND(' + "+".join(kcol(i, colk, pp) for i in range(KIDS)) + ',2),"")', fmt=NUM)
    yl = yt + ROWS
    notes = yl + 2
    put(ws, f"B{notes}", f'=IF({lump_c}>0.5,"The opening fund in "&PlanYear&"/"&RIGHT(PlanYear+1,2)&" is the amount needed now ("&RepCcy&" "&TEXT({lump_c},"#,##0")&"): costs that start before any yearly saving can be made. The plan assumes no existing savings.","The plan assumes no existing savings, so the fund starts at zero.")', MUTED)
    put(ws, f"B{notes + 1}", "Each year: opening fund - education expenses (paid at the start of the year) + investment growth + family savings "
                             "(added at the end of the year) = closing fund, which becomes the next year's opening fund.", MUTED)
    saved, grown = f"SUM(E{yt + 1}:E{yl})", f"SUM(F{yt + 1}:F{yl})"
    final = f"IF({hmax_c}>0,INDEX(H{yt + 1}:H{yl},{hmax_c}),0)"
    put(ws, f"B{notes + 2}", f'="Adds up: "&RepCcy&" "&TEXT({lump_c},"#,##0")&" needed now + "&RepCcy&" "&TEXT({saved},"#,##0")&" family savings + "&RepCcy&" "&TEXT({grown},"#,##0")&" investment growth - "&RepCcy&" "&TEXT({total_c},"#,##0")&" education expenses = "&RepCcy&" "&TEXT(IF(ABS({final})<0.5,0,{final}),"#,##0")&" left at the end."', MUTED)
    put(ws, f"I{notes}", None)
    # chart data (hidden columns), #N/A beyond the plan so the chart leaves gaps
    ch0 = 20  # column T
    put(ws, f"{L(ch0)}{yt}", "Year")
    for i in range(KIDS):
        put(ws, f"{L(ch0 + 1 + i)}{yt}", f"={kname(i)}")
    put(ws, f"{L(ch0 + 1 + KIDS)}{yt}", "Family savings")
    for pp in range(1, CHART_ROWS + 1):
        r = yt + pp
        show = f"{pp}<={hmax_c}"
        put(ws, f"{L(ch0)}{r}", f'=IF({show},B{r},"")')
        for i in range(KIDS):
            put(ws, f"{L(ch0 + 1 + i)}{r}", f"=IF({show},{kcol(i, 'D', pp)},NA())")
        put(ws, f"{L(ch0 + 1 + KIDS)}{r}", f"=IF({show},E{r},NA())")
    for c in range(ch0, ch0 + KIDS + 2):
        ws.column_dimensions[L(c)].hidden = True
    bar = BarChart()
    bar.type, bar.grouping, bar.overlap = "col", "stacked", 100
    bar.title = "Education costs by child and family savings, by year"
    bar.add_data(Reference(ws, min_col=ch0 + 1, max_col=ch0 + KIDS, min_row=yt, max_row=yt + CHART_ROWS), titles_from_data=True)
    bar.set_categories(Reference(ws, min_col=ch0, min_row=yt + 1, max_row=yt + CHART_ROWS))
    line = LineChart()
    line.add_data(Reference(ws, min_col=ch0 + 1 + KIDS, min_row=yt, max_row=yt + CHART_ROWS), titles_from_data=True)
    line.series[0].smooth = False
    bar += line
    bar.x_axis.delete = False
    bar.y_axis.delete = False
    bar.y_axis.numFmt = '#,##0'
    bar.y_axis.majorGridlines = None
    bar.height, bar.width = 9, 22
    bar.visible_cells_only = False   # the chart data sits in hidden columns
    ws.add_chart(bar, "J3")
    ws.freeze_panes = "A4"

    # ---------------- Compare Countries ----------------
    ws_c = ws_cc
    sel = lambda key: "CHOOSE(CmpChild," + ",".join(f"{q(S_IN)}!${L(3 + i)}${IN[key]}" for i in range(KIDS)) + ")"
    CB = []
    top = 3
    put(ws_c, "A1", "Compare Calc - the selected child's course in each country (helper sheet)", TITLE)
    ws_c.column_dimensions["A"].width = 44
    for c in countries:
        put(ws_c, f"A{top - 1}", c, H2)
        o = cost_block(ws_c, top, {"active": "IF(CmpChild<=NumChildren,1,0)", "country": f'"{c}"', "qual": sel("qualification"),
                                   "age": sel("age"), "entry": sel("entryAge"), "tinf": sel("tInf"), "linf": sel("lInf"), "ov": None})
        CB.append(o)
        top = o["end"] + 4
    ws_cc.sheet_state = "hidden"

    ws = ws_cmp
    ws.column_dimensions["A"].width = 2
    for c, w in {"B": 18, "C": 20, "D": 54, "E": 16}.items():
        ws.column_dimensions[c].width = w
    put(ws, "B1", "Compare countries", TITLE)
    put(ws, "B2", '="The same child and qualification in each country, in "&RepCcy&". Choose the child on the Inputs sheet (Compare countries for child number)."', MUTED)
    ch = lambda field: "CHOOSE(CmpChild," + ",".join(kref(i, KO[i][field]) for i in range(KIDS)) + ")"
    put(ws, "B4", "Child", BOLD)
    put(ws, "C4", "=CHOOSE(CmpChild," + ",".join(kname(i) for i in range(KIDS)) + ")")
    put(ws, "B5", "Qualification", BOLD)
    put(ws, "C5", f"={ch('qual')}")
    put(ws, "B6", "Current plan", BOLD)
    put(ws, "C6", f'=IF({ch("ok")}=1,{ch("country")},"Enter this child\'s details on the Inputs sheet first.")')
    header_row(ws, 8, 2, ["Country", '="Total ("&RepCcy&")"', "Note", "Current plan"])
    for j, c in enumerate(countries):
        r = 9 + j
        o = CB[j]
        is_cur = f'{ch("country")}="{c}"'
        put(ws, f"B{r}", c, BOLD)
        put(ws, f"C{r}", f'=IF({ch("ok")}<>1,"",IF({is_cur},{ch("total")},IF({q(S_CC)}!{o["ok"]}=1,{q(S_CC)}!{o["total"]},"Not available")))', BOLD, fmt=NUM)
        st, hm = f'IF({is_cur},{ch("tuition_status")},{q(S_CC)}!{o["tuition_status"]})', f'IF({is_cur},{ch("has_missing")},{q(S_CC)}!{o["has_missing"]})'
        put(ws, f"D{r}", f'=IF({ch("ok")}<>1,"",IF(AND(NOT({is_cur}),{q(S_CC)}!{o["ok"]}<>1),{q(S_CC)}!{o["msg"]},'
                         f'IF({st}="estimated","Estimate — no published fee for this course",IF({hm}=1,"Some costs have no figure yet",""))))', MUTED)
        put(ws, f"E{r}", f'=IF(AND({ch("ok")}=1,{is_cur}),"Current plan","")', BOLD)
    put(ws, "B16", "Each total covers the full course in that country (its usual length), at future prices, after your scholarship / "
                   "part-time-work percentage. The current plan's row is your plan exactly. For other countries your own fee and "
                   "living-cost increases are used, but your own cost figures are not carried over, because they are in another "
                   "country's currency.", MUTED, wrap=True)
    ws.merge_cells("B16:E16")
    ws.row_dimensions[16].height = 52

    # ---------------- Start Here ----------------
    ws = ws_start
    ws.column_dimensions["A"].width = 2
    ws.column_dimensions["B"].width = 110
    lines = [
        ("Children's Future Education Fund Calculator", TITLE),
        ("Sindhi Connect - sindhiconnect.org", MUTED),
        ("", None),
        ("How to use", H2),
        ("1. Go to the Inputs sheet and fill in the yellow cells: your family details and, for each child, age, education country and qualification.", None),
        ("2. Open the Dashboard: total education fund required, required yearly savings, amount needed now, each child, cost breakdown and the year-by-year plan.", None),
        ("3. Compare Countries shows the same course in all six study countries.", None),
        ("4. Each child's sheet shows every cost line with its status (Published figure / Estimated — please review / Not needed / Your figure) and the full calculation.", None),
        ("5. To use your own figures, type them in the 'Your own cost figures' rows on the Inputs sheet (today's prices, education country's currency).", None),
        ("", None),
        ("Important", H2),
        (d.get("disclaimer", ""), None),
        ("Results are in US dollars unless you choose another currency. The workbook uses formulas only (no macros). "
         "It follows the same method as the calculator on the website; see the Sources and Method sheet.", None),
        ("Fees checked on " + d["datasetDate"] + ". Exchange rates of " + d["exchangeRates"]["date"] + " (you can enter your own on the Assumptions and Currency sheet).", None),
    ]
    for i, (text, f) in enumerate(lines):
        c = put(ws, f"B{2 + i}", text, f or font(), wrap=True)
    ws.sheet_view.showGridLines = False

    # ---------------- Sources and Method ----------------
    ws = ws_src
    ws.column_dimensions["A"].width = 130
    method = [
        ("Method (same as the website calculator - docs/education-calculator/METHODOLOGY.md)", H2),
        ("Fee status: local fees when the child's nationality is a domestic nationality of the study country, otherwise international fees and a student visa.", None),
        ("University fees: the average of published fee records for the same country, qualification, fee status and currency, each grown to the plan start year by the fee increase. Where none exists, a labelled planning estimate is used.", None),
        ("Courses entered after another degree (e.g. medicine in the USA) use the pre-stage degree's fees for the first years.", None),
        ("Each cost line grows by its increase from its price year: line x (1 + increase)^(year - price year) x part of the year x exchange rate. One-time costs fall in the first study year.", None),
        ("Scholarship / part-time work reduces tuition, university fees and living costs once; visa and travel are not reduced.", None),
        ("Timing: costs are paid at the start of each year; savings are added at the end. The fund earns the investment return on what is left after costs.", None),
        ("Amount needed now: costs that come before any yearly saving can be made. Required yearly saving: the smallest first-year saving (rising by the yearly increase) that keeps the fund from running out in any year.", None),
        ("Yearly plan: opening fund - expenses + growth + savings = closing fund; each closing fund is the next opening fund. The plan assumes no existing savings.", None),
        ("", None),
        ("Sources", H2),
        ("Every published fee record used, with its source link, is on the Fee Records sheet. Planning-estimate bases are in the last column of the Cost Database sheet. "
         "Living costs, visas and travel are on their own sheets. The full research register is docs/education-calculator/SOURCES.md in the project.", None),
    ]
    for i, (text, f) in enumerate(method):
        put(ws, f"A{1 + i}", text, f or font(), wrap=True)

    # ---------------- protection ----------------
    for w in wb.worksheets:
        if w.title not in (S_LISTS, S_CC):
            w.protection.sheet = True
            w.protection.formatCells = False
            w.protection.formatColumns = False
            w.protection.formatRows = False
    wb.active = 1
    wb.calculation.fullCalcOnLoad = True
    return wb


def main():
    with open(DATA, encoding="utf-8") as f:
        d = json.load(f)
    wb = build_xlsx(d)
    os.makedirs(os.path.dirname(OUT_XLSX), exist_ok=True)
    wb.save(OUT_XLSX)
    print("Wrote", os.path.relpath(OUT_XLSX, ROOT))


if __name__ == "__main__":
    main()
